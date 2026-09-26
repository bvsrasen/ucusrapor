"""Karşılaştırma sonuçlarını tek bir Excel dosyasına yazar."""
from datetime import datetime

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

YAZI = "Calibri"
BASLIK = Font(name=YAZI, size=14, bold=True)
ALT = Font(name=YAZI, size=11, bold=True, color="1F4E79")
KALIN = Font(name=YAZI, size=10, bold=True)
NORMAL = Font(name=YAZI, size=10)
NOT = Font(name=YAZI, size=9, italic=True, color="666666")
GRI = PatternFill("solid", fgColor="E7E6E6")
ACIK = PatternFill("solid", fgColor="DDEBF7")
ince = Side(style="thin", color="BFBFBF")
CERCEVE = Border(left=ince, right=ince, top=ince, bottom=ince)


def _tablo(ws, satir, sutun, basliklar, veriler, formatlar=None):
    for j, b in enumerate(basliklar):
        c = ws.cell(row=satir, column=sutun + j, value=b)
        c.font, c.fill, c.border = KALIN, GRI, CERCEVE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, satir_verisi in enumerate(veriler, start=1):
        for j, deger in enumerate(satir_verisi):
            c = ws.cell(row=satir + i, column=sutun + j, value=deger)
            c.font, c.border = NORMAL, CERCEVE
            if formatlar and formatlar[j]:
                c.number_format = formatlar[j]
    return satir + len(veriler)


def _grafik(ws_veri, x_sutun, y_sutunlar, ilk, son, baslik, x_ad, y_ad, renkler):
    g = ScatterChart()
    g.title, g.style, g.height, g.width = baslik, 13, 8, 18
    g.x_axis.title, g.y_axis.title = x_ad, y_ad
    g.x_axis.delete = g.y_axis.delete = False
    x = Reference(ws_veri, min_col=x_sutun, min_row=ilk, max_row=son)
    for sutun, renk in zip(y_sutunlar, renkler):
        s = Series(Reference(ws_veri, min_col=sutun, min_row=ilk - 1, max_row=son), x, title_from_data=True)
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = renk
        s.graphicalProperties.line.width = 15000
        g.series.append(s)
    return g


def _veri_sayfasi(wb, ad, basliklar, sutunlar):
    ws = wb.create_sheet(ad)
    for j, b in enumerate(basliklar, start=1):
        c = ws.cell(row=1, column=j, value=b)
        c.font, c.fill = KALIN, GRI
        ws.column_dimensions[get_column_letter(j)].width = 14
    for i, degerler in enumerate(zip(*sutunlar), start=2):
        for j, v in enumerate(degerler, start=1):
            ws.cell(row=i, column=j, value=None if v is None or (isinstance(v, float) and np.isnan(v)) else round(float(v), 3))
    ws.freeze_panes = "A2"
    return ws, len(sutunlar[0]) + 1


def rapor_yaz(cikti, masa, vakumlar, ucus, ateslemeler, kaynak_klasor):
    wb = Workbook()
    oz = wb.active
    oz.title = "Özet"
    oz.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [34, 14, 14, 14, 14, 16, 3, 12]):
        oz.column_dimensions[col].width = w

    oz["A1"] = "UKB Test Karşılaştırma Raporu"
    oz["A1"].font = BASLIK
    oz["A2"] = f"Veri klasörü: {kaynak_klasor}   |   Oluşturulma: {datetime.now():%d.%m.%Y %H:%M}"
    oz["A2"].font = NOT
    oz["A3"] = "Not: Bu rapordaki veriler örnek (simüle) veridir, gerçek test kaydı değildir."
    oz["A3"].font = Font(name=YAZI, size=9, italic=True, color="C00000")

    # --- masa testi
    r = 5
    oz.cell(row=r, column=1, value="1. Masa testi (10 dk, durağan)").font = ALT
    r = _tablo(oz, r + 1, 1, ["", "Özgün UKB", "Ticari UKB"], [
        ["İrtifa gürültüsü (std, m)", masa["ozgun_gurultu_m"], masa["ticari_gurultu_m"]],
        ["10 dk'daki kayma (m)", masa["ozgun_kayma_m"], masa["ticari_kayma_m"]],
        ["Kayıp örnek (%)", masa["ozgun_kayip_yuzde"], None],
        ["Kart sıcaklık artışı (°C)", masa["sicaklik_artisi_c"], None],
    ], [None, "0.00", "0.00"])
    oz.cell(row=r + 1, column=1,
            value="Özgün kartın kayması kart ısındıkça basınç sensörünün kaymasından kaynaklanıyor.").font = NOT

    # --- vakum testi
    r += 3
    oz.cell(row=r, column=1, value="2. Vakum odası testi – tepe noktası tespiti (s)").font = ALT
    bas = r + 1
    satirlar = []
    for k, v in enumerate(vakumlar):
        rr = bas + 1 + k
        satirlar.append([f"Deneme {v['deneme']}", v["vana_acilis_s"], v["ozgun_basit_s"], v["ozgun_filtreli_s"],
                         v["ticari_s"], f"=E{rr}-B{rr}"])
    r = _tablo(oz, bas, 1, ["", "Vana açıldı (referans)", "Özgün – mevcut algoritma", "Özgün – filtreli",
                           "Ticari", "Ticari gecikme"], satirlar, [None, "0.00", "0.00", "0.00", "0.00", "+0.00;-0.00"])
    ort = r + 1
    oz.cell(row=ort, column=1, value="Ortalama fark (tespit − referans)").font = KALIN
    for col in "CDE":
        c = oz[f"{col}{ort}"]
        c.value = f"=AVERAGE({col}{bas + 1}:{col}{r})-AVERAGE(B{bas + 1}:B{r})"
        c.number_format = "+0.00;-0.00"
        c.font, c.border, c.fill = KALIN, CERCEVE, ACIK
    oz.cell(row=ort + 1, column=1,
            value="Negatif = erken tetikleme. Pompanın oluşturduğu basınç dalgalanması mevcut algoritmayı yanıltıyor.").font = NOT
    oz.cell(row=ort + 2, column=1,
            value=f"Kartların saatleri senkron değil; ticari kayıt irtifa eğrileri üst üste getirilerek "
                  f"{', '.join(str(v['saat_farki_s']) for v in vakumlar)} s kaydırıldı.").font = NOT

    # --- uçuş simülasyonu
    r = ort + 4
    oz.cell(row=r, column=1, value="3. Uçuş simülasyonu – tepe noktası yöntemleri").font = ALT
    oz.cell(row=r + 1, column=1, value="Gerçek tepe noktası (kalkıştan sonra, s)").font = KALIN
    ref = oz.cell(row=r + 1, column=2, value=ucus["gercek_tepe_s"])
    ref.number_format, ref.border = "0.00", CERCEVE
    ref_adres = f"$B${r + 1}"
    bas = r + 3
    satirlar = []
    for k, s in enumerate(ucus["tablo"].itertuples(index=False)):
        rr = bas + 1 + k
        satirlar.append([s.yontem, s.tespit_s, s.tespit_irtifa_m, f"=B{rr}-{ref_adres}",
                         f'=IF(ABS(D{rr})<=1,"uygun",IF(D{rr}<0,"erken","geç"))'])
    r = _tablo(oz, bas, 1, ["Yöntem", "Tespit (s)", "Tespit irtifası (m)", "Hata (s)", "Sonuç"],
               satirlar, [None, "0.00", "#,##0", "+0.00;-0.00", None])
    notlar = [
        "Mevcut algoritma ses hızı civarındaki basınç hatası yüzünden ~1300 m'de, roket hâlâ çok hızlıyken tetikliyor.",
        "İvmeölçer ±16 g'de doyuma girdiği için ivmeden hesaplanan hız düşük çıkıyor; Mach kilidi erken açılıyor.",
        "Zamanlayıcı süresi simülasyondan alındı; gerçek uçuşta sapma olabilir, tek başına değil yedek olarak düşünülmeli.",
    ]
    for k, n in enumerate(notlar):
        oz.cell(row=r + 1 + k, column=1, value=n).font = NOT

    # --- ateşleme
    r += len(notlar) + 2
    oz.cell(row=r, column=1, value="4. Yer ateşleme testleri").font = ALT
    oz.cell(row=r + 1, column=1, value="Toplam deneme").font = KALIN
    oz.cell(row=r + 1, column=2, value="=COUNTA('Ateşleme'!A2:A100)")
    oz.cell(row=r + 2, column=1, value="Başarılı ayrılma").font = KALIN
    oz.cell(row=r + 2, column=2, value="=COUNTIF('Ateşleme'!F2:F100,\"evet\")")
    oz.cell(row=r + 3, column=1, value="Süreklilik sorunu").font = KALIN
    oz.cell(row=r + 3, column=2, value="=COUNTIF('Ateşleme'!D2:D100,0)")
    r += 5

    # --- çıkarımlar
    oz.cell(row=r, column=1, value="Karar için çıkarımlar").font = ALT
    cikarimlar = [
        "1) Özgün kartın mevcut tepe noktası algoritması bu hâliyle kurtarmayı tetiklemek için güvenilir değil.",
        "2) Özgün kartta zaman kilidi + filtreli barometrik yöntem denenmeli; aynı testler tekrar edilip kaydedilmeli.",
        "3) Ticari kart vakum testlerinde tutarlı (~0,3 s gecikme) ama tek başına bırakılmamalı; iki kart karşılaştırılarak ana/yedek kararı verilmeli.",
        "4) Rampada, hakem altimetresi takıldıktan SONRA iki kartın süreklilik ve yer istasyonu bağlantısı tekrar kontrol edilmeli.",
    ]
    for k, c in enumerate(cikarimlar):
        oz.cell(row=r + 1 + k, column=1, value=c).font = NORMAL

    # --- veri sayfaları ve grafikler
    v1 = vakumlar[0]
    t_o = v1["ozgun"]["t_s"].to_numpy()
    h_t = np.interp(t_o, v1["ticari"]["zaman_hizali_s"], v1["ticari"]["irtifa_m"], left=np.nan, right=np.nan)
    ws_v, son_v = _veri_sayfasi(wb, "Vakum 1", ["Zaman (s)", "Özgün (m)", "Ticari (m)"],
                                [t_o, v1["ozgun"]["irtifa_m"].to_numpy(), h_t])
    oz.add_chart(_grafik(ws_v, 1, [2, 3], 2, son_v, "Vakum denemesi 1 – irtifa eşdeğeri", "Zaman (s)",
                         "İrtifa (m)", ["2E75B6", "ED7D31"]), "H5")

    u = ucus["df"]
    ws_u, son_u = _veri_sayfasi(wb, "Uçuş Sim", ["Kalkıştan sonra (s)", "İrtifa (m)", "Eksenel ivme (g)"],
                                [u["t_s"].to_numpy() - ucus["kalkis_s"], u["irtifa_m"].to_numpy(), u["az_g"].to_numpy()])
    oz.add_chart(_grafik(ws_u, 1, [2], 2, son_u, "Uçuş simülasyonu – özgün UKB barometrik irtifa",
                         "Kalkıştan sonra (s)", "İrtifa (m)", ["2E75B6"]), "H22")

    m = masa["ozgun"].iloc[::10]
    mt = np.interp(m["t_s"], masa["ticari"]["zaman_s"], masa["ticari"]["irtifa_m"])
    ws_m, son_m = _veri_sayfasi(wb, "Masa", ["Zaman (s)", "Özgün (m)", "Ticari (m)", "Kart sıcaklığı (°C)"],
                                [m["t_s"].to_numpy(), m["irtifa_m"].to_numpy(), mt, m["sicaklik_c"].to_numpy()])
    oz.add_chart(_grafik(ws_m, 1, [2, 3], 2, son_m, "Masa testi – durağan irtifa okuması", "Zaman (s)",
                         "İrtifa (m)", ["2E75B6", "ED7D31"]), "H39")

    def sayi(x):
        try:
            return float(x) if "." in x else int(x)
        except ValueError:
            return x if x != "" else None

    ws_a = wb.create_sheet("Ateşleme")
    _tablo(ws_a, 1, 1, ["Deneme", "UKB", "Barut (g)", "Süreklilik (Ω)", "Gecikme (ms)", "Ayrılma", "Not"],
           [[sayi(str(x)) for x in satir] for satir in ateslemeler.values.tolist()])
    for col, w in zip("ABCDEFG", [8, 8, 10, 13, 12, 9, 60]):
        ws_a.column_dimensions[col].width = w

    ws_k = wb.create_sheet("Veri Kalitesi")
    tum = []
    for kaynak, kayit in [("Masa", masa["kayit"])] + [(f"Vakum {v['deneme']}", v["kayit"]) for v in vakumlar] + \
                         [("Uçuş Sim", ucus["kayit"])]:
        tum += [[kaynak, k["tip"], k["t_ms"], k["detay"], k["islem"]] for k in kayit]
    _tablo(ws_k, 1, 1, ["Kayıt", "Sorun", "Zaman (ms)", "Detay", "İşlem"], tum)
    for col, w in zip("ABCDE", [11, 22, 11, 30, 30]):
        ws_k.column_dimensions[col].width = w
    ws_k.freeze_panes = "A2"
    ws_k.auto_filter.ref = f"A1:E{len(tum) + 1}"

    wb.save(cikti)
    return cikti
