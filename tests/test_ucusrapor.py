"""
UçuşRapor otomatik test senaryoları (UR-18).

Her test, Confluence'taki "05 - Test Senaryoları" sayfasındaki bir senaryoya
(TS-xx) ve gereksinim dokümanındaki bir gereksinime karşılık gelir.

Çalıştırma:  python -m pytest -v
"""

import json
import time
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from ucusrapor.isleme import (
    VeriFormatHatasi, metrikleri_hesapla, temizle, veri_oku, zenginlestir,
)
from ucusrapor.rapor import excel_raporu_olustur, rapor_bilgisi

KOK = Path(__file__).resolve().parent.parent
LOG = KOK / "veri" / "ucus_log_simule.csv"
BEKLENEN = KOK / "veri" / "beklenen_hatalar.json"


@pytest.fixture(scope="module")
def ham():
    return veri_oku(LOG)


@pytest.fixture(scope="module")
def sonuc(ham):
    return temizle(ham)


@pytest.fixture(scope="module")
def beklenen():
    with open(BEKLENEN, encoding="utf-8") as f:
        return json.load(f)


def _anahtarlar(kayitlar, tip):
    return {(k["zaman_ms"], k["sutun"]) for k in kayitlar if k["tip"] == tip}


# --- TS-01 / FR-01 ---------------------------------------------------------
def test_ts01_log_dosyasi_okunur(ham):
    assert len(ham) > 1000
    assert {"zaman_ms", "irtifa_m", "durum"}.issubset(ham.columns)


def test_ts01b_eksik_sutunlu_dosya_reddedilir(tmp_path):
    bozuk = tmp_path / "bozuk.csv"
    pd.DataFrame({"zaman_ms": [0, 100], "irtifa_m": [0, 1]}).to_csv(bozuk, index=False)
    with pytest.raises(VeriFormatHatasi):
        veri_oku(bozuk)


def test_ts01c_olmayan_dosya_anlasilir_hata_verir():
    with pytest.raises(FileNotFoundError):
        veri_oku("olmayan_dosya.csv")


# --- TS-02 / FR-02 ---------------------------------------------------------
def test_ts02_tum_eksik_degerler_tespit_edilir(sonuc, beklenen):
    beklenen_set = _anahtarlar(beklenen, "EKSIK_DEGER")
    bulunan_set = _anahtarlar(sonuc.kayitlar, "EKSIK_DEGER")
    assert beklenen_set == bulunan_set


def test_ts02b_temiz_veride_bos_hucre_kalmaz(sonuc):
    assert not sonuc.temiz.isna().any().any()


# --- TS-03 / FR-03 ---------------------------------------------------------
def test_ts03_duplike_satirlar_silinir(sonuc, beklenen):
    assert _anahtarlar(beklenen, "DUPLIKE_SATIR") == _anahtarlar(sonuc.kayitlar, "DUPLIKE_SATIR")
    assert sonuc.temiz["zaman_ms"].is_unique


# --- TS-04 / FR-04 ---------------------------------------------------------
def test_ts04_sensor_sicramalari_yakalanir(sonuc, beklenen):
    assert _anahtarlar(beklenen, "SENSOR_SICRAMASI") == _anahtarlar(sonuc.kayitlar, "SENSOR_SICRAMASI")


def test_ts04b_yanlis_alarm_uretilmez(sonuc, beklenen):
    """Motor ateşlemesi gibi gerçek ani değişimler hata sayılmamalı."""
    assert len(sonuc.kayitlar) == len(beklenen)


def test_ts04c_gecersiz_durum_kodu_silinir(sonuc):
    assert set(sonuc.temiz["durum"].unique()) <= set(range(8))


# --- TS-05 / FR-05 ---------------------------------------------------------
def test_ts05_metrikler_elle_hesaplananla_ortusur(sonuc):
    df = zenginlestir(sonuc.temiz)
    m = metrikleri_hesapla(df)

    # "Elle" hesap: doğrudan pandas ile bağımsız yoldan
    assert m["maks_irtifa_m"] == pytest.approx(df["irtifa_m"].max())
    assert m["ucus_suresi_s"] == pytest.approx(
        df.loc[df.durum == 7, "ucus_zamani_s"].iloc[0] - df.loc[df.durum == 1, "ucus_zamani_s"].iloc[0])
    assert m["maks_ivme_g"] == pytest.approx(
        np.sqrt(df.ivme_x_g**2 + df.ivme_y_g**2 + df.ivme_z_g**2).max())


def test_ts05b_metrikler_fiziksel_olarak_tutarli(sonuc):
    m = metrikleri_hesapla(zenginlestir(sonuc.temiz))
    assert 0 < m["apogee_zamani_s"] < m["suruklenme_parasutu_s"] < m["ana_parasut_s"] < m["inis_zamani_s"]
    assert m["suruklenme_inis_hizi_ms"] > m["ana_parasut_inis_hizi_ms"] > 0
    assert m["ana_parasut_irtifa_m"] < m["maks_irtifa_m"]


# --- TS-06 / FR-06, FR-07, FR-08, NFR-01 ----------------------------------
def test_ts06_excel_raporu_tek_adimda_ve_hizli_uretilir(sonuc, tmp_path):
    bas = time.perf_counter()
    df = zenginlestir(sonuc.temiz)
    m = metrikleri_hesapla(df)
    cikti = tmp_path / "rapor.xlsx"
    excel_raporu_olustur(df, sonuc, m, cikti, rapor_bilgisi("test.csv", 0.0, True))
    sure = time.perf_counter() - bas

    assert cikti.exists()
    assert sure < 300, "NFR-01: rapor 5 dakikanın altında üretilmeli"

    wb = load_workbook(cikti)
    assert wb.sheetnames == ["Özet", "Temiz Veri", "Veri Kalitesi", "Uçuş Olayları"]
    assert len(wb["Özet"]._charts) == 3, "İrtifa, ivme ve veri kalitesi grafikleri olmalı"
    assert wb["Veri Kalitesi"].max_row - 4 == len(sonuc.kayitlar)
